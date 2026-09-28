from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


EXPORT_FIELDS = [
    ("频道名称", "channel_name"), ("频道 URL", "channel_url"), ("频道 ID", "channel_id"), ("地区判定", "region"),
    ("地区置信度", "region_confidence"), ("判定依据", "region_basis"), ("内容标签", "tags"), ("命中关键词", "matched_keywords"),
    ("订阅数", "subscribers"), ("总视频数", "total_videos"), ("近90天发片数", "posts_90"), ("近期视频平均观看量", "avg_views"),
    ("互动率估算", "engagement_rate"), ("公开联系入口", "contacts"), ("潮流勁抽推广识别", "promotion_status"),
    ("检索时间", "checked_at"), ("检索视频数量", "scanned_videos"), ("已推广证据", "evidence"), ("合作匹配度", "score"),
    ("评分说明", "score_reason"), ("人工审核状态", "review_status"), ("备注", "note"), ("收藏", "favorite"),
    ("未成年人风险", "minor_risk"), ("数据来源", "source"), ("采集时间", "sourced_at"), ("最近5条相关视频", "videos"),
    ("合规提示", "compliance_notice"),
]
NOTICE = "需确认当地广告披露、抽奖及未成年人营销合规；“未发现推广过潮流勁抽”不等于从未推广。"


def _format(value, key: str) -> str | int | float:
    if key == "compliance_notice":
        return NOTICE
    if key == "engagement_rate":
        return round(float(value or 0) * 100, 2)
    if key in {"tags", "matched_keywords", "contacts"}:
        return "；".join(value or [])
    if key == "evidence":
        return " | ".join(f"{item.get('title', '')} {item.get('url', '')} 命中:{'、'.join(item.get('hits', []))}" for item in (value or []))
    if key == "videos":
        return " | ".join(f"{item.get('title', '')} [{item.get('published_at', '')}] {item.get('views', 0)}次 {item.get('url', '')}" for item in (value or []))
    if key in {"favorite", "minor_risk"}:
        return "是" if value else "否"
    return value if value is not None else ""


def export_rows(profiles: list[dict]) -> list[list]:
    return [[_format(profile.get(key), key) for _, key in EXPORT_FIELDS] for profile in profiles]


def export_csv(path: str | Path, profiles: list[dict]) -> Path:
    path = Path(path)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([label for label, _ in EXPORT_FIELDS])
        writer.writerows(export_rows(profiles))
    return path


def export_xlsx(path: str | Path, profiles: list[dict]) -> Path:
    path = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "达人名单"
    sheet.append([label for label, _ in EXPORT_FIELDS])
    for row in export_rows(profiles):
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, (label, key) in enumerate(EXPORT_FIELDS, 1):
        width = 18
        if key in {"region_basis", "evidence", "score_reason", "videos", "compliance_notice"}:
            width = 38
        elif key in {"channel_url", "source"}:
            width = 30
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=2, min_col=13, max_col=13):
        row[0].number_format = '0.00"%"'
    sheet.row_dimensions[1].height = 28
    workbook.save(path)
    return path
